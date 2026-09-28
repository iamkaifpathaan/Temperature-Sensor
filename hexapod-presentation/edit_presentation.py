import copy, re, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

FIG = "/home/user/Temperature-Sensor/hexapod-progress-report/figures/"
SRC = "/home/user/Temperature-Sensor/hexapod-progress-report/photos/src/"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"a": A}
FONT = '<a:latin typeface="Times New Roman" panose="02020603050405020304"/><a:ea typeface="Times New Roman" panose="02020603050405020304"/><a:cs typeface="Times New Roman" panose="02020603050405020304"/><a:sym typeface="Times New Roman" panose="02020603050405020304"/>'
NAVY = "002F4A"
STYLE_ID = "{6B23DE0B-3846-4328-826E-2A21398C07C0}"

p = Presentation(sys.argv[1])
S = list(p.slides)


def esc(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def runs_xml(text, size, color=None, italic=False, bold_all=False):
    out = []
    for part in re.split(r"(\*\*.+?\*\*|__.+?__)", text):
        if not part:
            continue
        b = bold_all or part.startswith("**")
        it = italic or part.startswith("__")
        t = part[2:-2] if part.startswith(("**", "__")) else part
        col = f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>' if color else '<a:solidFill><a:schemeClr val="tx1"/></a:solidFill>'
        out.append(f'<a:r><a:rPr lang="en-IN" sz="{int(size*100)}" b="{1 if b else 0}" i="{1 if it else 0}" dirty="0">{col}{FONT}</a:rPr><a:t xml:space="preserve">{esc(t)}</a:t></a:r>')
    return "".join(out)


def para_xml(text, size=18, bullet=True, after=8, align="l", color=None, italic=False, bold=False, hang=0.3, num=None):
    if bullet or num:
        m = int(Inches(hang))
        bu = (f'<a:buClr><a:srgbClr val="{NAVY}"/></a:buClr><a:buSzPct val="100000"/><a:buFont typeface="Arial"/><a:buChar char="•"/>' if bullet
              else f'<a:buClr><a:schemeClr val="tx1"/></a:buClr><a:buSzPct val="100000"/><a:buFont typeface="+mj-lt"/><a:buAutoNum type="{num}"/>')
        ppr = f'<a:pPr marL="{m}" indent="-{m}" algn="{align}"><a:lnSpc><a:spcPct val="100000"/></a:lnSpc><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="{after*100}"/></a:spcAft>{bu}</a:pPr>'
    else:
        ppr = f'<a:pPr marL="0" indent="0" algn="{align}"><a:lnSpc><a:spcPct val="100000"/></a:lnSpc><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="{after*100}"/></a:spcAft><a:buNone/></a:pPr>'
    return f'<a:p xmlns:a="{A}">{ppr}{runs_xml(text, size, color, italic, bold)}</a:p>'


def fill_text(shape, paras, anchor="t", autofit_off=True):
    """paras: list of xml strings from para_xml."""
    tb = shape.text_frame._txBody
    for p_ in tb.findall(f"{{{A}}}p"):
        tb.remove(p_)
    for x in paras:
        tb.append(etree.fromstring(x))
    bp = tb.find(f"{{{A}}}bodyPr")
    for c in list(bp):
        bp.remove(c)
    bp.set("wrap", "square"); bp.set("anchor", anchor)
    for k in ("lIns", "rIns"):
        bp.set(k, "91425")


def geom(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def textbox(slide, x, y, w, h, paras, anchor="t"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(tb, paras, anchor)
    return tb


def title_shape(slide):
    best = None
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            for r in sh._element.iter(f"{{{A}}}rPr"):
                if r.get("sz") in ("2800", "4000") and (r.get("b") == "1" or r.get("sz") == "4000"):
                    return sh
    return best


def set_title(slide, text):
    sh = title_shape(slide)
    fill_text(sh, [f'<a:p xmlns:a="{A}"><a:pPr marL="0" indent="0" algn="ctr"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>'
                   f'<a:r><a:rPr lang="en-IN" sz="2800" b="1" dirty="0"><a:solidFill><a:schemeClr val="lt1"/></a:solidFill>{FONT}</a:rPr><a:t>{esc(text)}</a:t></a:r></a:p>'], anchor="ctr")
    geom(sh, 1.2, 0.5, 10.93, 0.7)
    return sh


def body_shape(slide, exclude):
    for sh in slide.shapes:
        if sh.has_text_frame and sh is not exclude and sh.shape_id != exclude.shape_id and sh.text_frame.text.strip():
            return sh
    return None


def remove(shape):
    shape._element.getparent().remove(shape._element)


def picture(slide, path, x, y, w=None, h=None, caption=None, cap_w=None):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    ln = pic.line; ln.color.rgb = RGBColor(0x00, 0x2F, 0x4A); ln.width = Pt(0.75)
    if caption:
        cw = cap_w or pic.width / 914400
        cx = pic.left / 914400 + (pic.width / 914400 - cw) / 2
        textbox(slide, cx, (pic.top + pic.height) / 914400 + 0.05, cw, 0.35,
                [para_xml(caption, 12, bullet=False, after=0, align="ctr", italic=True)])
    return pic


# ---------- tables
def cell_text(cell, text, size=12, bold=False, color=None, align="l"):
    tb = cell._tc.find(f"{{{A}}}txBody")
    for p_ in tb.findall(f"{{{A}}}p"):
        tb.remove(p_)
    tb.append(etree.fromstring(para_xml(text, size, bullet=False, after=0, align=align, color=color, bold=bold)))


def new_table(slide, x, y, widths, header, rows, size=13, hsize=13, row_h=0.42, colors=None, bold_col=None, hrow=0.55):
    shp = slide.shapes.add_table(len(rows) + 1, len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * (len(rows) + 1)))
    tbl = shp.table
    sid = tbl._tbl.tblPr.find(f"{{{A}}}tableStyleId")
    if sid is None:
        sid = etree.SubElement(tbl._tbl.tblPr, f"{{{A}}}tableStyleId")
    sid.text = STYLE_ID
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for j, h in enumerate(header):
        cell_text(tbl.cell(0, j), h, hsize, bold=True, color="FFFFFF")
    for i, row in enumerate(rows, 1):
        tbl.rows[i].height = Inches(row_h)
        for j, v in enumerate(row):
            c = colors(i - 1, j) if colors else None
            cell_text(tbl.cell(i, j), v, size, color=c, bold=(bool(c) and j > 0) or j == bold_col)
    tbl.rows[0].height = Inches(hrow)
    return shp


# ======================================================================
# Slide 1 — only fix the double space in the title
for sh in S[0].shapes:
    if sh.has_text_frame:
        for t in sh._element.iter(f"{{{A}}}t"):
            if t.text and "EMERGENCY  AND" in t.text:
                t.text = t.text.replace("EMERGENCY  AND", "EMERGENCY AND")

# Slide 3 — POs: keep PO1–PO5 only (user's wording), tidy PSO indentation
s = S[2]
box = [sh for sh in s.shapes if sh.has_text_frame and "POs" in sh.text_frame.text][0]
paras = box.text_frame._txBody.findall(f"{{{A}}}p")
texts = ["".join(t.text or "" for t in q.iter(f"{{{A}}}t")) for q in paras]
keep_until = next(i for i, t in enumerate(texts) if t.strip().startswith("POs"))
head = paras[:keep_until + 1]
pso_head = next(q for q, t in zip(paras, texts) if t.strip().startswith("PSOs"))
tb = box.text_frame._txBody
for q in paras:
    tb.remove(q)
for q in head:
    tb.append(q)


def labeled(label, text, size=18):
    m = int(Inches(0.75))
    return (f'<a:p xmlns:a="{A}"><a:pPr marL="{m}" indent="-{m}" algn="l"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="500"/></a:spcAft>'
            f'<a:buNone/><a:tabLst><a:tab pos="{m}" algn="l"/></a:tabLst></a:pPr>'
            f'<a:r><a:rPr lang="en-IN" sz="{size*100}" b="1" dirty="0"><a:solidFill><a:srgbClr val="{NAVY}"/></a:solidFill>{FONT}</a:rPr><a:t>{label}</a:t></a:r>'
            f'<a:r><a:rPr lang="en-IN" sz="{size*100}" b="0" dirty="0"><a:solidFill><a:schemeClr val="tx1"/></a:solidFill>{FONT}</a:rPr><a:t xml:space="preserve">\t{esc(text)}</a:t></a:r></a:p>')


POS = [("PO1:", "Application of engineering knowledge in electronics, embedded systems, sensors and mechanical design."),
       ("PO2:", "Analysis of obstacle detection, stability, torque, current and power requirements."),
       ("PO3:", "Design and development of an autonomous hexapod surveillance robot with integrated sensing."),
       ("PO4:", "Testing and investigation of servo, sensor, IMU, power and system performance."),
       ("PO5:", "Use of Arduino, ESP32-CAM, PCA9685, HC-SR04, MPU6050 and relevant engineering tools.")]
for l, t in POS:
    tb.append(etree.fromstring(labeled(l, t)))
tb.append(pso_head)
tb.append(etree.fromstring(labeled("PSO1:", "Graduates will be able to design, analyze, and implement electronic circuits, and use techniques relevant to electronics and communication engineering.")))
tb.append(etree.fromstring(labeled("PSO2:", "Graduates will demonstrate the ability to apply communication engineering principles in developing innovative and sustainable solutions in areas such as wireless communication systems and antenna design.")))

# Slide 4 — TOC
s = S[3]
t = title_shape(s); body = body_shape(s, t)
toc = ["Introduction & Motivation", "Problem Statement & Objectives", "Literature Review", "Base Paper Findings",
       "Methodology", "System Architecture", "Hardware Components & Tools Used", "Planned Implementation",
       "Design Comparison", "Custom PCB Development", "Current Project Status", "Future Plan of Action", "References (IEEE Format)"]
fill_text(body, [para_xml(x, 20, bullet=False, num="arabicPeriod", after=5, hang=0.5) for x in toc])
geom(body, 3.6, 1.35, 6.8, 5.9)
set_title(s, "TABLE OF CONTENTS")

# Slide 5 — Introduction (planned wording)
s = S[4]; t = set_title(s, "INTRODUCTION"); body = body_shape(s, t)
fill_text(body, [para_xml(x, 20, after=10) for x in [
    "**Disaster response need:** search-and-rescue teams often have to inspect collapsed, unstable or contaminated areas that are unsafe for people.",
    "**Why six legs:** a hexapod can keep at least three feet on the ground while walking, so it stays statically stable and can step over rubble and gaps where wheeled robots get stuck.",
    "**Controller:** the robot is designed around an Arduino Uno R3, which will run the gait algorithm, read the sensors and command every joint.",
    "**Actuation:** 18 MG996R metal-gear servos (three per leg), driven through two PCA9685 16-channel PWM drivers on a shared I²C bus.",
    "**Sensing:** an MPU6050 IMU for body tilt and three HC-SR04 ultrasonic sensors (front, left, right) for obstacle detection.",
    "**Surveillance:** an ESP32-CAM on a pan–tilt mount is planned to stream live video over Wi-Fi.",
    "**Wiring:** a custom PCB shield replaces roughly 90 loose jumper wires — its design is complete and the Gerber files are ready."]])
geom(body, 0.6, 1.4, 12.1, 5.8)

# Slide 6 — Motivation
s = S[5]; t = set_title(s, "MOTIVATION"); body = body_shape(s, t)
fill_text(body, [para_xml(x, 20, after=16) for x in [
    "**Uneven environments:** wheeled robots struggle with debris, stairs and uneven ground typical of disaster zones and industrial sites.",
    "**Operator safety:** sending people into hazardous or physically inaccessible areas carries high risk.",
    "**Compact embedded platform:** need for a small, self-contained robot that manages its own power and complex locomotion.",
    "**Remote visual monitoring:** live video lets a team assess a situation safely from a distance.",
    "**Affordable and reproducible:** commercial inspection and rescue robots are too expensive for student labs; building from off-the-shelf parts keeps the project practical and repeatable."]])
geom(body, 0.6, 1.45, 12.1, 5.6)

# Slide 7 — Problem statement
s = S[6]; t = set_title(s, "PROBLEM STATEMENT"); body = body_shape(s, t)
fill_text(body, [para_xml(x, 22, bullet=False, num="arabicPeriod", after=18, hang=0.5) for x in [
    "Multi-legged locomotion with 3 degrees of freedom (DOF) per leg for better terrain adaptability.",
    "Embedded control of 18 heavily loaded servo motors at the same time without signal degradation.",
    "Stable, multi-rail power distribution that isolates logic components from high-current actuator spikes.",
    "Remote visual surveillance through a pan–tilt camera over Wi-Fi.",
    "Reliable wiring on a machine that vibrates with every step."]])
geom(body, 0.8, 1.6, 11.7, 5.2)

# Slide 8 — Objectives (content kept)
s = S[7]; t = set_title(s, "OBJECTIVES"); body = body_shape(s, t)
fill_text(body, [para_xml(x, 22, bullet=False, num="arabicPeriod", after=22, hang=0.5) for x in [
    "To design and develop a low-cost six-legged (hexapod) robot using an Arduino Uno and servo motors that can walk stably over flat and moderately uneven surfaces.",
    "To integrate ultrasonic distance sensing and an ESP32-CAM module so that the robot can perform basic obstacle avoidance and stream live video for surveillance.",
    "To design and evaluate a control scheme that coordinates the leg movements and sensor feedback for safe navigation in emergency scenarios."]])
geom(body, 0.8, 1.7, 11.7, 4.8)

# Slides 9–11 — Literature review (verified references only)
LIT = [
    ("1", "Ji et al. — IEEE Access, 2024 (Base paper)", "RL-based gait and foot-trajectory planning; static stability margin used as the stability measure; tilt-based gait switching.", "Needs RL training (PyBullet) and a GPU — not feasible on an Arduino. We adopt only the stability concepts."),
    ("2", "Mohamed et al. — Proc. IMechE Part C, 2024", "Kinematic model and body levelling of a hexapod with telescopic legs.", "Telescopic actuators exceed our budget; the modelling approach guided our leg torque check."),
    ("3", "Liu et al. — J. Field Robotics, 2025", "Crab-like lateral walking with hierarchical control on unknown rugged terrain.", "Long lateral walking loads the hip joints heavily; we start with a forward tripod gait."),
    ("4", "Xu et al. — J. Field Robotics, 2024", "Improved knee joint raises obstacle clearance without extra actuators.", "Shows the effect of leg geometry; our knee uses a standard rotary servo."),
    ("5", "Saengsint et al. — IEEE ROBIO, 2023", "Rescue hexapod with AI-based human detection and tracking.", "On-board AI needs far more processing than a microcontroller."),
    ("6", "Kulkarni et al. — Springer LNEE, 2024", "Arduino Uno + ultrasonic sensor + ESP32-CAM surveillance robot with live video.", "Wheeled platform; confirms the sensing and camera stack we are using."),
    ("7", "Mahmud et al. — IEEE Access, 2024", "Synthetic images used to improve victim detection in disaster scenes.", "Shows the value of a visual link; victim detection is outside our scope."),
    ("8", "Cruz Ulloa et al. — Applied Sciences, 2024", "Compares thermal, multispectral and RGB vision for victim detection.", "Thermal / multispectral cameras exceed our budget; RGB (ESP32-CAM) chosen."),
    ("9", "Jadeja et al. — Scientific Reports, 2024", "Deep-learning survivor detection on a snake robot.", "Needs neural-network hardware; detection is better done on the operator side."),
    ("10", "Gelfert — RiTA 2022 (Springer), 2023", "Real-time victim detection in smoke using a multi-sensor unit.", "Specialised sensors and deep learning; beyond this project."),
]
groups = [LIT[0:4], LIT[4:7], LIT[7:10]]
for k, (slide, rows) in enumerate(zip(S[8:11], groups)):
    for sh in list(slide.shapes):
        if sh.has_text_frame and ("Literature" in sh.text_frame.text or "comprehensive" in sh.text_frame.text):
            remove(sh)
        elif sh.has_table:
            remove(sh)
    tb_ = textbox(slide, 1.2, 0.5, 10.93, 0.7, [para_xml("LITERATURE REVIEW" + ("" if k == 0 else " (CONTD.)"), 28, bullet=False, after=0, align="ctr", color="FFFFFF", bold=True)], anchor="ctr")
    new_table(slide, 0.45, 1.45, [0.6, 2.75, 4.6, 4.45], ["S. No.", "REFERENCE", "FINDINGS", "LIMITATION / RELEVANCE TO OUR PROJECT"],
              [list(r) for r in rows], size=16, hsize=16, row_h=1.15 if len(rows) == 4 else 1.25, bold_col=1)

# Slide 12 — Base paper findings
s = S[11]; t = set_title(s, "BASE PAPER FINDINGS"); body = body_shape(s, t)
fill_text(body, [para_xml("**Base paper:** S. Ji, W. Wei, X. Liu and J. Wu, “Optimization Control of Attitude Stability for Hexapod Robots Based on Reinforcement Learning,” IEEE Access, vol. 12, 2024.", 18, bullet=False, after=14)] +
          [para_xml(x, 20, after=11) for x in [
              "**Static stability:** the robot is statically stable when the projection of its centre of mass (CoM) lies inside the support polygon formed by the feet on the ground.",
              "**Stability margin (SSM):** the shortest distance from the CoM projection to the polygon edges — a larger margin means safer walking on rough ground.",
              "**Hierarchical RL control (RL-FI-RW):** a policy network outputs only four gait parameters (X, Y, lift height H, rotation Ω); inverse kinematics and a trajectory controller drive the joints.",
              "**Tilt-based gait switching:** tripod gait normally, 4-legged gait above 20° body inclination and 5-legged gait at 25°.",
              "**Results (simulation only):** in PyBullet, with a 15 kg robot model, the method showed less body sway and longer travel than CPG-based methods on rough terrain."]])
geom(body, 0.6, 1.4, 12.1, 5.9)

# Slide 13 — Base paper: relevance to our project (+ support polygon diagram)
s = S[12]; t = set_title(s, "BASE PAPER — RELEVANCE TO OUR PROJECT"); body = body_shape(s, t)
fill_text(body, [para_xml("**What we use from the paper**", 18, bullet=False, after=6, color=NAVY)] +
          [para_xml(x, 18, after=8) for x in [
              "Support polygon and stability margin → tripod gait, and a low, central centre of mass (battery and converters under the body plate).",
              "Body-inclination feedback → MPU6050 IMU in our design.",
              "More legs on the ground for rough terrain → planned 4-legs-down gait.",
              "Smooth swing trajectories instead of abrupt joint moves."]] +
          [para_xml("**What we do not use**", 18, bullet=False, after=6, color=NAVY)] +
          [para_xml(x, 18, after=8) for x in [
              "The reinforcement-learning controller — it needs GPU training and cannot run on an Arduino Uno (2 KB RAM).",
              "Research gap addressed: applying the stability concepts to a low-cost, microcontroller-based physical robot."]])
geom(body, 0.5, 1.4, 6.9, 5.9)
picture(s, FIG + "fig_support_polygon.png", 7.55, 2.1, w=5.35, caption="Support polygon and stability margin during a tripod gait")

# Slide 14 — Methodology
s = S[13]; t = set_title(s, "METHODOLOGY"); body = body_shape(s, t)
fill_text(body, [para_xml(x, 20, after=12) for x in [
    "**Phase 1 — Architecture & PCB design:** multi-rail power distribution finalised; custom PCB shield designed in KiCad.  __(Completed)__",
    "**Phase 2 — Mechanical fabrication:** 3D printing of the legs, body plate, electronics box and camera mount.  __(Next)__",
    "**Phase 3 — Assembly & integration:** PCB soldering; servos, sensors and ESP32-CAM integrated step by step.  __(Planned)__",
    "**Phase 4 — Kinematics & gait:** inverse kinematics for 18 joints on the Arduino Uno; tripod gait based on the static stability margin.  __(Planned)__",
    "**Phase 5 — Command loop:** web page → ESP32-CAM serial link → Arduino gait engine → PCA9685 (I²C) → MG996R servos.  __(Planned)__",
    "**Phase 6 — Safety & testing:** stop forward motion if an obstacle is within 30 cm; tilt and low-battery checks; calibration and testing.  __(Planned)__"]])
geom(body, 0.6, 1.4, 12.1, 5.9)

# Slide 15 — System architecture
s = S[14]; t = set_title(s, "SYSTEM ARCHITECTURE"); body = body_shape(s, t)
picture(s, FIG + "fig_block_diagram.png", 0.55, 1.35, h=5.55, caption="Planned system architecture of the hexapod")
fill_text(body, [para_xml(x, 18, after=12) for x in [
    "**Arduino Uno R3** is the only controller: gait, sensor reading and joint commands.",
    "**ESP32-CAM** is a camera peripheral: video goes to the phone over Wi-Fi; single-character commands reach the Uno over a serial link.",
    "**Power:** 3S LiPo → 20 A fuse → switch → four buck converters.",
    "**Rails:** two 6 V servo rails (XL4016), a 5 V logic rail and a separate 5 V camera rail (LM2596).",
    "**One custom PCB shield** carries all power zones, driver modules and connectors, with a single star-ground point."]])
geom(body, 7.35, 1.55, 5.45, 5.6)

# Slide 16 — Hardware components & tools
s = S[15]; t = set_title(s, "HARDWARE COMPONENTS & TOOLS USED"); body = body_shape(s, t)
fill_text(body, [para_xml(x, 18, after=11) for x in [
    "**Controller & camera:** Arduino Uno R3; ESP32-CAM (OV2640).",
    "**Actuation:** 18 × MG996R (legs), 2 × SG90 (camera pan/tilt), 2 × PCA9685 PWM drivers.",
    "**Sensing:** 3 × HC-SR04 ultrasonic sensors, MPU6050 IMU.",
    "**Power:** 3S LiPo 2200 mAh (two packs), 2 × XL4016 (6 V), 2 × LM2596 (5 V), 20 A fuse and switch.",
    "**Software:** KiCad, Arduino IDE, VS Code, Git/GitHub.",
    "**Cost:** estimated purchase cost ≈ ₹17,346 (build-plan estimate); 3D printing provided by the college.",
    "**Status:** all required components have been procured; none installed yet."]])
geom(body, 0.5, 1.4, 6.9, 5.9)
picture(s, FIG + "photo_components.jpg", 7.6, 1.4, h=5.3, caption="Some of the procured components")

# Slides 17–18 — Proposed implementation (planned)
IMPL = [
    ("PLANNED IMPLEMENTATION — CONTROL", [
        "**Inverse kinematics:** foot position (x, y, z) converted into coxa–femur–tibia joint angles (θ1, θ2, θ3) for each leg.",
        "**Tripod gait:** the six legs moved in two alternating groups of three (swing and stance).",
        "**PWM offloading:** two PCA9685 boards (I²C addresses 0x40 and 0x41) generate the 50 Hz servo pulses in hardware.",
        "**IMU posture correction:** MPU6050 pitch and roll used to correct the body posture on slopes.",
        "**Ultrasonic interlock:** front, left and right HC-SR04 sensors checked in the main loop to override forward motion.",
        "**Calibration & memory:** servos centred (PWM count ≈ 307) before fitting horns; gait tables stored in PROGMEM to save the Uno's 2 KB RAM."]),
    ("PLANNED IMPLEMENTATION — POWER & HARDWARE", [
        "**Multi-rail power:** XL4016 converters for the servo rails and LM2596 for logic, keeping servo current away from the logic circuits.",
        "**Camera supply:** a dedicated converter and ≥ 1000 µF capacitors to absorb the ESP32-CAM's ~700 mA Wi-Fi bursts.",
        "**Servo master-disable:** PCA9685 OE line pulled HIGH at boot, so servos stay off until a safe pose is written.",
        "**Battery monitoring:** resistor divider on A0; warning at 10.5 V and safe sit-down at 9.9 V.",
        "**Star ground:** all negatives meet at one point to avoid ground loops and I²C errors.",
        "**Mechanical:** converters in an open, slotted tray under the body plate for cooling; brass heat-set inserts in 3D-printed parts."]),
]
for s, (title, items) in zip(S[16:18], IMPL):
    t = set_title(s, title); body = body_shape(s, t)
    fill_text(body, [para_xml(x, 20, after=13) for x in items])
    geom(body, 0.6, 1.4, 12.1, 5.3)
    textbox(s, 0.6, 6.8, 12.1, 0.4, [para_xml("These features are part of the planned design and will be verified during integration and testing.", 14, bullet=False, after=0, italic=True)])

# Slide 19 — Design comparison
s = S[18]; t = set_title(s, "DESIGN COMPARISON: CONVENTIONAL vs PROPOSED"); body = body_shape(s, t); remove(body)
new_table(s, 0.6, 1.45, [2.1, 4.7, 5.3], ["ASPECT", "TYPICAL HOBBY BUILD", "OUR DESIGN APPROACH"], [
    ["Wiring", "Breadboard and ~90 loose jumper wires", "Custom PCB shield with labelled sockets"],
    ["Servo PWM", "Pulses generated by the Arduino", "Two PCA9685 boards generate pulses in hardware"],
    ["Servo power", "Single 5 V rail, often from the Arduino", "Two 6 V XL4016 rails through PCB copper zones"],
    ["Camera power", "Shares the logic supply", "Dedicated converter with bulk capacitors"],
    ["Protection", "Often none", "20 A fuse, rocker switch, OE interlock, battery sensing"],
    ["Cooling", "Regulators inside a closed box", "Open, slotted converter tray under the body plate"],
    ["Mobility", "Wheels stop at rubble and gaps", "18-DOF legs; thigh load ≈ 47 % of servo stall torque (design estimate)"]],
    size=17, hsize=17, row_h=0.66)
textbox(s, 0.6, 6.85, 12.1, 0.5, [para_xml("Expected benefits are based on design calculations and will be confirmed through testing after assembly.", 14, bullet=False, after=0, italic=True)])

# Slide 20 — PCB development (1/2): schematic and layout
s = S[19]; t = set_title(s, "CUSTOM PCB DEVELOPMENT (1/2)"); body = body_shape(s, t)
fill_text(body, [para_xml("Schematic ✓   →   PCB Layout ✓   →   3D View ✓   →   Gerber Generation ✓   →   Fabrication: Pending", 17, bullet=False, after=0, align="ctr", color=NAVY, bold=True)])
geom(body, 0.6, 1.2, 12.1, 0.45)
picture(s, FIG + "kicad_schematic_1.png", 0.55, 1.85, w=6.3, caption="Schematic in KiCad (sheet 1 of 2)")
picture(s, FIG + "kicad_layout_2d.png", 7.2, 1.85, h=3.9, caption="PCB layout in KiCad")
textbox(s, 0.6, 5.3, 6.3, 2.0, [para_xml(x, 16, after=6) for x in [
    "Shield for the Arduino Uno, about 114 × 84 mm, two copper layers.",
    "Separate zones: 6 V zone A (right legs + camera servos), 6 V zone B (left legs), 5 V logic and an isolated 5 V camera node.",
    "20 labelled servo sockets, sensor headers and PCA9685 module sockets."]])

# Slide 21 — PCB development (2/2): 3D view, checks, Gerbers
s = S[20]; t = set_title(s, "CUSTOM PCB DEVELOPMENT (2/2)"); body = body_shape(s, t)
picture(s, FIG + "kicad_3d_top.png", 0.55, 1.4, h=5.2, caption="3D view of the PCB in KiCad")
fill_text(body, [para_xml("**Design checks**", 18, bullet=False, after=4, color=NAVY)] +
          [para_xml(x, 16, after=4) for x in ["ERC: 0 errors, 0 warnings.", "DRC: 0 unconnected items; remaining markers are silkscreen clearance (printed labels only)."]] +
          [para_xml("**Gerber files generated (27 Sep 2026)**", 18, bullet=False, after=4, color=NAVY)] +
          [para_xml(x, 16, after=4) for x in ["Front and back copper, solder mask, silkscreen and paste layers.", "Board outline (Edge Cuts), PTH and NPTH drill files, Gerber job file."]] +
          [para_xml("**Status**", 18, bullet=False, after=4, color=NAVY),
           para_xml("PCB design complete and Gerber files ready. **Fabrication / ordering is the next pending step** — the board has not been manufactured yet.", 16, bullet=False, after=0)])
geom(body, 7.55, 1.45, 5.3, 5.6)

# Slide 22 — Current project status
s = S[21]; t = set_title(s, "CURRENT PROJECT STATUS"); body = body_shape(s, t); remove(body)
STAT = [["Planning & research", "✓ Completed"], ["Component selection & procurement", "✓ Completed"],
        ["PCB schematic & layout", "✓ Completed"], ["Design checks (ERC / DRC)", "✓ Completed"],
        ["Gerber file generation", "✓ Completed"], ["PCB ordering / fabrication", "Pending — next step"],
        ["3D printing of mechanical parts", "Next"], ["Mechanical assembly & PCB installation", "Planned"],
        ["Electronics, sensor & camera integration", "Planned"], ["Programming, calibration & testing", "Planned"]]


def scol(i, j):
    if j != 1:
        return None
    v = STAT[i][1]
    return "1B5E20" if v.startswith("✓") else ("9A3412" if v.startswith(("Pending", "Next")) else "37474F")


new_table(s, 0.6, 1.45, [4.9, 2.7], ["STAGE", "STATUS"], STAT, size=17, hsize=17, row_h=0.48, colors=scol)
picture(s, SRC + "11.png", 8.4, 2.2, w=4.45, caption="PCB design — ready for fabrication")


# Slide 23 — Future plan of action
s = S[22]; t = set_title(s, "FUTURE PLAN OF ACTION")
tshape = [sh for sh in s.shapes if sh.has_table][0]; tbl = tshape.table
PLAN = [["Month", "Activity", "Status"],
        ["AUG", "Literature review, base-paper study and architecture finalisation", "Completed"],
        ["SEP", "Component procurement; PCB schematic, layout, design checks and Gerber generation", "Completed"],
        ["OCT", "Order and fabricate the PCB; start 3D printing of legs and body parts", "To be done"],
        ["NOV", "Assemble the mechanical structure; mount servos and electronics; install the PCB", "To be done"],
        ["DEC", "Integrate sensors and ESP32-CAM; develop gait control; calibrate the system", "To be done"],
        ["JAN", "Test walking, stability, obstacle sensing and surveillance; final demonstration", "To be done"]]
for i, row in enumerate(PLAN):
    for j, v in enumerate(row):
        col = "FFFFFF" if i == 0 else ("1B5E20" if v == "Completed" else ("9A3412" if v == "To be done" else None))
        cell_text(tbl.cell(i, j), v, 18, bold=(i == 0 or j == 0 or j == 2), color=col, align="l")
    tbl.rows[i].height = Inches(0.62)
for j, w in enumerate([1.1, 8.6, 2.1]):
    tbl.columns[j].width = Inches(w)
tshape.left, tshape.top = Inches(0.75), Inches(1.5)
textbox(s, 0.75, 6.3, 11.8, 0.6, [para_xml("Plan is tentative and depends on PCB delivery time and 3D-printer availability.", 14, bullet=False, after=0, italic=True)])

# Slides 24–25 — References (verified, corrected)
REFS = [
    'S. Ji, W. Wei, X. Liu, and J. Wu, “Optimization control of attitude stability for hexapod robots based on reinforcement learning,” IEEE Access, vol. 12, pp. 154120–154132, 2024.',
    'S. Mohamed, H. Q. Le, Y. Kim, and B. Shin, “Design and modeling of hexapod robot using telescopic legs connected to pivot joints at the hips,” Proc. IMechE, Part C: J. Mech. Eng. Sci., vol. 238, 2024.',
    'C. Liu, Y. Zhu, Z. Han, et al., “Design and control of a hexapod robot RENS H3 for lateral walking on unknown rugged terrains,” J. Field Robotics, 2025.',
    'K. Xu, R. Qin, C. Chen, G. Dong, J. Chen, and X. Ding, “Design and multimodal locomotion plan of a hexapod robot with improved knee joints,” J. Field Robotics, vol. 41, 2024.',
    'C. Saengsint et al., “Autonomous rescue hexapod robot with AI human detection and tracking,” in Proc. IEEE Int. Conf. Robotics and Biomimetics (ROBIO), Koh Samui, Thailand, 2023.',
    'S. V. Kulkarni, N. Samanvita, S. Gatade, and R. Likhitha, “Obstacle avoidance robot based on Arduino for live video transmission and surveillance,” Lecture Notes in Electrical Engineering, vol. 1104, Springer, 2024.',
    'S. Mahmud, A. A. Fime, and J.-H. Kim, “ATR HarmoniSAR: A system for enhancing victim detection in robot-assisted disaster scenarios,” IEEE Access, vol. 12, 2024.',
    'C. Cruz Ulloa, D. Orbea, J. del Cerro, and A. Barrientos, “Thermal, multispectral, and RGB vision systems analysis for victim detection in SAR robotics,” Applied Sciences, vol. 14, no. 2, p. 766, 2024.',
    'R. Jadeja, T. Trivedi, and J. Surve, “Survivor detection approach for post earthquake search and rescue missions based on deep learning inspired algorithms,” Scientific Reports, vol. 14, 2024.',
    'S. Gelfert, “Real time victim detection in smoky environments with mobile robot and multi-sensor unit using deep learning,” in Robot Intelligence Technology and Applications 7 (RiTA 2022), Springer, 2023.',
]


def ref_para(n, text):
    m = int(Inches(0.55))
    return (f'<a:p xmlns:a="{A}"><a:pPr marL="{m}" indent="-{m}" algn="l"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="1000"/></a:spcAft><a:buNone/><a:tabLst><a:tab pos="{m}" algn="l"/></a:tabLst></a:pPr>'
            + runs_xml(f"[{n}]\t" + text, 19) + '</a:p>')


for k, s in enumerate(S[23:25]):
    t = set_title(s, "REFERENCES" + ("" if k == 0 else " (CONTD.)")); body = body_shape(s, t)
    fill_text(body, [ref_para(5 * k + i + 1, r) for i, r in enumerate(REFS[5 * k:5 * k + 5])])
    geom(body, 0.7, 1.45, 11.9, 5.8)

# Slide 26 — Thank you: remove leading spaces, centre
s = S[25]
for sh in s.shapes:
    if sh.has_text_frame and "THANK" in sh.text_frame.text:
        for tt in sh._element.iter(f"{{{A}}}t"):
            tt.text = tt.text.strip()
        for pp in sh._element.iter(f"{{{A}}}pPr"):
            pp.set("algn", "ctr")
        geom(sh, 2.5, 3.0, 8.33, 1.1)

p.save(sys.argv[2])
print("saved")
